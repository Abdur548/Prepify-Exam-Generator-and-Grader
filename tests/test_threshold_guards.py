"""Guards that a mutation audit found untested (Problem 3, 2026-09-01).

## What the audit did

Every guard that compares against a `config` threshold was disabled in turn — the
comparison replaced so it can never fire — and the suite re-run. A guard whose
removal breaks nothing is a guard no test exercises: the code around it passes
whether or not it works, and moving its threshold would go unnoticed.

Nine guards, seven caught, two not:

    RELEVANCE_FLOOR              caught  test_a_span_item_is_still_evaluated_normally
    DEDUP_TAU                    caught  test_duplication_gate_flags_similar_items
    TOPIC_MATCH_MIN_EVIDENCE     caught  test_evidence_floor_rejects_a_match_on_a_ubi...
    TOPIC_MATCH_RELATIVE_FLOOR   caught  test_items_come_only_from_matched_nodes
    SYNTHESIS_ITEM_WARN_RATIO    caught  test_a_paper_over_the_ratio_warns
    COGNITIVE_BALANCE_TOLERANCE  caught  test_divergence_warns_and_does_not_raise
    EQUATION_MIN_MATH_CHARS      caught  test_math_unicode_flagged
    OPTION_LENGTH_BAND           NOT CAUGHT
    RERANKER_THRESHOLD           NOT CAUGHT

This file covers the two that were not.

## Why the fixtures here are derived, not written down

The original concern was fixtures that hardcode a value chosen relative to a
threshold: change the threshold and the test silently stops testing anything. That
happened once already — a `-1.0` scorer fixture was a rejection at `TAU=3.5` and
became a pass at `RELEVANCE_FLOOR=-2.0`, going green while asserting nothing.

So every value below is computed FROM the constant it exercises. If someone moves
`RERANKER_THRESHOLD` or `OPTION_LENGTH_BAND`, these tests move with it and keep
testing the boundary rather than a number that used to be near it.
"""
from __future__ import annotations

import pytest

from coursegen import config


# ---------------------------------------------------------------------------
# RERANKER_THRESHOLD — via should_use_material()
# ---------------------------------------------------------------------------

def _reranked(score: float):
    from coursegen.retrieve.rerank import RerankedChunk
    return RerankedChunk(chunk_id="c1", text="some text", file="deck.pdf", page=1, score=score)


class TestRerankerThresholdDefault:
    """`should_use_material` decides `from_material`, and its DEFAULT was untested.

    The existing tests in test_p4_render_chat.py pass `threshold=3.5` and
    `threshold=9.0` explicitly, so they never touch
    `config.RERANKER_THRESHOLD` — the value production actually uses. Changing that
    constant could not fail a single test.

    This matters beyond tidiness. `should_use_material` is what sets
    `from_material` on a chat answer. If it wrongly returns True, the UI tells a
    student the answer came from their own material when nothing relevant was
    retrieved.
    """

    def test_default_threshold_is_read_from_config(self) -> None:
        from coursegen.retrieve.rerank import should_use_material
        # Derived from the constant: comfortably either side, whatever it is set to.
        above = _reranked(config.RERANKER_THRESHOLD + 1.0)
        below = _reranked(config.RERANKER_THRESHOLD - 1.0)

        assert should_use_material([above]) is True
        assert should_use_material([below]) is False

    def test_score_exactly_at_the_threshold_is_accepted(self) -> None:
        """The comparison is `>=`. Pinning it stops a silent flip to `>`."""
        from coursegen.retrieve.rerank import should_use_material
        assert should_use_material([_reranked(config.RERANKER_THRESHOLD)]) is True

    def test_no_chunks_means_not_from_material(self) -> None:
        """Empty retrieval must not read as 'grounded in your material'."""
        from coursegen.retrieve.rerank import should_use_material
        assert should_use_material([]) is False

    def test_explicit_threshold_still_overrides_the_default(self) -> None:
        from coursegen.retrieve.rerank import should_use_material
        chunk = _reranked(config.RERANKER_THRESHOLD + 1.0)
        assert should_use_material([chunk], threshold=chunk.score + 1.0) is False


# ---------------------------------------------------------------------------
# OPTION_LENGTH_BAND — via _mcq_hygiene_issue()
# ---------------------------------------------------------------------------

def _outlier_length(base_len: int, n_options: int) -> int:
    """Smallest outlier length that pushes the short options below the band.

    Solved rather than guessed, so the test tracks the constant:

        min_allowed = avg * (1 - band),  avg = ((n-1)*L + K) / n

    the short options fall outside when  L < avg * (1 - band), i.e.

        K > L * (n - (1 - band)(n - 1)) / (1 - band)

    A hardcoded "make one option 100 chars" would stop triggering if the band
    widened, and the test would pass while checking nothing — the exact failure
    this file exists to prevent.
    """
    band = config.OPTION_LENGTH_BAND
    if band >= 1.0:  # a band that wide admits everything; no outlier exists
        pytest.skip(f"OPTION_LENGTH_BAND={band} admits any length")
    k = base_len * (n_options - (1 - band) * (n_options - 1)) / (1 - band)
    return int(k) + 2  # strictly greater, plus a margin for integer rounding


def _item_with_option_lengths(lengths: list[int]):
    from coursegen.contracts.item import GeneratedItem, MCQOption, SourceRef
    return GeneratedItem(
        slot_id="A-01",
        stem="Which of the following is correct?",
        options=[
            # A DISTINCT character per option. Repeating one character makes two
            # equal-length options identical strings, and the duplicate-option
            # check fires first — the fixture would then be testing the wrong
            # guard while appearing to test this one.
            MCQOption(label=chr(ord("A") + i), text=chr(ord("a") + i) * n)
            for i, n in enumerate(lengths)
        ],
        correct_option="A",
        model_answer="A",
        explanation="because",
        source_ref=SourceRef(file="deck.pdf", pages=[1]),
    )


class TestOptionLengthBand:
    """The MCQ option-length band had NO test at all.

    Disabling it entirely — `min_allowed = 0`, `max_allowed = inf` — broke nothing
    in the suite. The gate exists because a conspicuously longer option is a
    giveaway: students answer it without reading the question, and the item stops
    measuring anything.
    """

    def test_options_outside_the_band_are_rejected(self) -> None:
        from coursegen.exam.validate import _mcq_hygiene_issue
        lengths = [20, 20, 20, _outlier_length(20, 4)]
        issue = _mcq_hygiene_issue(_item_with_option_lengths(lengths))
        assert issue is not None, (
            f"lengths {lengths} should breach OPTION_LENGTH_BAND="
            f"{config.OPTION_LENGTH_BAND} but were accepted"
        )
        assert "band" in issue.lower()

    def test_equal_length_options_are_accepted(self) -> None:
        """The other side of the boundary: the guard must not reject a clean item."""
        from coursegen.exam.validate import _mcq_hygiene_issue
        assert _mcq_hygiene_issue(_item_with_option_lengths([20, 20, 20, 20])) is None

    def test_variation_inside_the_band_is_accepted(self) -> None:
        """Real options are never identical in length; mild variation must pass.

        Sized from the constant at half the permitted deviation, so it stays
        comfortably inside whatever the band is set to.
        """
        base = 40
        delta = int(base * config.OPTION_LENGTH_BAND * 0.5)
        lengths = [base - delta, base, base, base + delta]
        assert _mcq_hygiene_issue_for(lengths) is None, (
            f"lengths {lengths} are within OPTION_LENGTH_BAND="
            f"{config.OPTION_LENGTH_BAND} but were rejected"
        )


def _mcq_hygiene_issue_for(lengths: list[int]):
    from coursegen.exam.validate import _mcq_hygiene_issue
    return _mcq_hygiene_issue(_item_with_option_lengths(lengths))


# ---------------------------------------------------------------------------
# Resource guards on third-party uploads — second audit batch
# ---------------------------------------------------------------------------
#
# The same mutation audit over the S3 resource limits found three more uncovered
# guards. These matter more than the quality thresholds above: they are the limits
# standing between an arbitrary uploaded file and this process. Disabling any of
# them broke nothing in the suite.
#
#     MAX_FILE_SIZE_BYTES            NOT CAUGHT
#     MAX_PAGES (pdf)                NOT CAUGHT
#     MAX_PAGES (pptx)               NOT CAUGHT
#     MAX_DECOMPRESSED_SIZE_BYTES    caught  (the zip-bomb guard is covered)
#     MIN_FREE_VIRTUAL_MEMORY_GB     caught
#
# `_parse_docx` is deliberately excluded: it documents why MAX_PAGES does NOT
# apply to a format with no pages, so there is no guard there to test.
#
# Each test patches the constant DOWN rather than generating a 100 MB file or a
# 2,001-page PDF. That keeps the suite fast, and — more importantly — the fixture
# is sized from the patched constant, so the test cannot go vacuous if the real
# limit changes.

class TestUploadResourceGuards:

    def test_oversized_file_is_rejected(self, tmp_path, monkeypatch) -> None:
        from coursegen.ingest.parse import _check_file_size
        monkeypatch.setattr(config, "MAX_FILE_SIZE_BYTES", 128)
        big = tmp_path / "big.pdf"
        big.write_bytes(b"x" * (config.MAX_FILE_SIZE_BYTES + 1))
        with pytest.raises(ValueError, match="MAX_FILE_SIZE_BYTES"):
            _check_file_size(big)

    def test_file_at_the_limit_is_accepted(self, tmp_path, monkeypatch) -> None:
        """The comparison is `>`, so exactly-at-the-limit must pass."""
        from coursegen.ingest.parse import _check_file_size
        monkeypatch.setattr(config, "MAX_FILE_SIZE_BYTES", 128)
        ok = tmp_path / "ok.pdf"
        ok.write_bytes(b"x" * config.MAX_FILE_SIZE_BYTES)
        _check_file_size(ok)  # must not raise

    def test_pdf_over_max_pages_is_rejected(self, tmp_path, monkeypatch) -> None:
        import fitz
        from coursegen.ingest.parse import _parse_pdf
        monkeypatch.setattr(config, "MAX_PAGES", 2)
        path = tmp_path / "long.pdf"
        doc = fitz.open()
        for i in range(config.MAX_PAGES + 1):
            doc.new_page().insert_text((72, 100), f"Page {i}", fontsize=11)
        doc.save(str(path)); doc.close()
        with pytest.raises(ValueError, match="MAX_PAGES"):
            _parse_pdf(path)

    def test_pdf_at_max_pages_is_accepted(self, tmp_path, monkeypatch) -> None:
        import fitz
        from coursegen.ingest.parse import _parse_pdf
        monkeypatch.setattr(config, "MAX_PAGES", 2)
        path = tmp_path / "short.pdf"
        doc = fitz.open()
        for i in range(config.MAX_PAGES):
            doc.new_page().insert_text((72, 100), f"Page {i}", fontsize=11)
        doc.save(str(path)); doc.close()
        _parse_pdf(path)  # must not raise

    def test_pptx_over_max_pages_is_rejected(self, tmp_path, monkeypatch) -> None:
        from pptx import Presentation
        from coursegen.ingest.parse import _parse_pptx
        monkeypatch.setattr(config, "MAX_PAGES", 2)
        path = tmp_path / "long.pptx"
        prs = Presentation()
        for _ in range(config.MAX_PAGES + 1):
            prs.slides.add_slide(prs.slide_layouts[5])
        prs.save(str(path))
        with pytest.raises(ValueError, match="MAX_PAGES"):
            _parse_pptx(path)
