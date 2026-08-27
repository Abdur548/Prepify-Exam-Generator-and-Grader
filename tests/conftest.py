"""Pytest configuration. The mocked LLM client is the default for all tests.
Live API tests require --live and are excluded from the default run (L3).
"""
from __future__ import annotations

from pathlib import Path

import pytest
from unittest.mock import MagicMock

from coursegen.llm.client import LLMClient, TokenBudget


def pytest_addoption(parser: pytest.Parser) -> None:
    parser.addoption(
        "--live",
        action="store_true",
        default=False,
        help="Run tests that make real LLM API calls (excluded by default).",
    )


def pytest_collection_modifyitems(
    config: pytest.Config, items: list[pytest.Item]
) -> None:
    if not config.getoption("--live"):
        skip_live = pytest.mark.skip(reason="Pass --live to run live API tests.")
        for item in items:
            if "live" in item.keywords:
                item.add_marker(skip_live)


@pytest.fixture
def mock_llm_client() -> MagicMock:
    """Default mocked LLM client. Replays a minimal valid response."""
    client = MagicMock(spec=LLMClient)
    client.budget = TokenBudget()
    client.call.return_value = {"mock": True, "estimated_tokens": 100}
    return client


@pytest.fixture
def token_budget() -> TokenBudget:
    return TokenBudget()


# ---------------------------------------------------------------------------
# P1 document fixtures — created once per test session via PyMuPDF / python-pptx.
# Using session scope so the heavy PDF creation is not repeated per test.
# ---------------------------------------------------------------------------

@pytest.fixture(scope="session")
def native_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    Native-text PDF: body at 11 pt, chapter headings at 24 pt.
    Font ratio is large enough that the relative threshold definitely fires.
    """
    import fitz  # PyMuPDF

    path = tmp_path_factory.mktemp("pdfs") / "native.pdf"
    doc = fitz.open()

    # Page 1 — Chapter 1
    p = doc.new_page()
    p.insert_text((72, 100), "Chapter 1: Introduction", fontsize=24)
    for y, line in enumerate([
        "This section introduces fundamental concepts.",
        "It covers the basic principles and key definitions.",
        "Understanding these concepts is essential for further study.",
        "We will explore various aspects in structured detail.",
        "Each concept builds carefully on the previous one.",
    ], start=1):
        p.insert_text((72, 100 + y * 60), line, fontsize=11)

    # Page 2 — Chapter 2
    p = doc.new_page()
    p.insert_text((72, 100), "Chapter 2: Data Structures", fontsize=24)
    for y, line in enumerate([
        "Data structures organise information efficiently in memory.",
        "Arrays provide fast indexed access to homogeneous elements.",
        "Linked lists allow dynamic allocation without contiguous memory.",
        "Trees represent hierarchical parent-child relationships.",
    ], start=1):
        p.insert_text((72, 100 + y * 60), line, fontsize=11)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def slide_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    Slide-exported PDF: body at 24 pt, slide titles at 36 pt.
    Regression fixture: an absolute font-size threshold would classify
    every 24 pt body block as a heading. The relative threshold must not.
    Gate: heading_fraction < HEADING_MAX_BLOCK_PCT (0.30).
    """
    import fitz

    path = tmp_path_factory.mktemp("pdfs") / "slides.pdf"
    doc = fitz.open()

    for i in range(5):
        p = doc.new_page()
        p.insert_text((72, 80), f"Lecture {i + 1}: Module Overview", fontsize=36)
        p.insert_text((72, 220), "First key point with supporting detail.", fontsize=24)
        p.insert_text((72, 320), "Second key point elaborated with an example.", fontsize=24)
        p.insert_text((72, 420), "Third key point summarising the concept.", fontsize=24)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def figure_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """PDF that contains at least one image block so has_figure can be verified."""
    import fitz

    path = tmp_path_factory.mktemp("pdfs") / "figure.pdf"
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Section with Diagram", fontsize=20)
    p.insert_text((72, 200), "Figure description text follows below.", fontsize=11)

    # Insert a 4×4 white rectangle as an image using PyMuPDF's own Pixmap.
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4))
    pix.clear_with(255)  # white
    img_rect = fitz.Rect(72, 250, 200, 350)
    p.insert_image(img_rect, pixmap=pix)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def malformed_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Garbage bytes with a .pdf extension — must fail gracefully."""
    path = tmp_path_factory.mktemp("pdfs") / "malformed.pdf"
    path.write_bytes(b"This is not a valid PDF file, just garbage bytes.")
    return path


@pytest.fixture(scope="session")
def minimal_pptx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Two-slide PPTX with title placeholders."""
    from pptx import Presentation

    path = tmp_path_factory.mktemp("pptx") / "test.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Introduction to Algorithms"
    slide.placeholders[1].text = "This slide covers fundamental algorithm concepts."

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Sorting Algorithms"
    slide.placeholders[1].text = "We discuss merge sort and quicksort in detail."

    prs.save(str(path))
    return path
