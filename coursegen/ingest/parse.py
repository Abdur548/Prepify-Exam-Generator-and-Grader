"""
Document parsing: PDF (PyMuPDF), PPTX and DOCX (python-pptx / python-docx)
→ TextBlock list.

Zero LLM calls. Zero heading decisions — that is structure.py's job.
This module extracts blocks, records their font metadata, and sets the per-block
content flags that coursemap.py ORs up into `NodeFlags`.

Content flags — what is a detector and what is a heuristic:

  has_table    DETECTED. PDF: `page.find_tables()`, present in the pinned PyMuPDF
               (1.27.2); a text block whose bbox intersects a detected table's bbox
               is marked. PPTX: `shape.has_table`, which is exact. DOCX: a `w:tbl`
               body element, which is exact.
  has_code     HEURISTIC. Monospace font name only (config.MONOSPACE_FONT_SUBSTRINGS).
               Prose set in a monospace face is a false positive; code set in a
               proportional face is a false negative. There is no lexical or
               syntactic analysis, and nothing in the pinned stack offers one.
  has_equation HEURISTIC. Either a dedicated mathematics font family
               (config.MATH_FONT_SUBSTRINGS) or at least
               config.EQUATION_MIN_MATH_CHARS characters drawn from
               config.MATH_UNICODE_RANGES in the block's text. Mathematics typeset
               as an *image* — common in scanned notes and exported slides — is
               invisible to both prongs, and MVP1 has no OCR, so it is a false
               negative. Both prongs are deliberately conservative; config.py records
               which font names and Unicode blocks were excluded and why.
"""
from __future__ import annotations

import logging
import zipfile
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FuturesTimeoutError
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from coursegen import config

logger = logging.getLogger(__name__)

# Suffixes parse_directory PICKS UP — not the same thing as "can be parsed".
# `.ppt` is scanned deliberately even though it can never be read: dropping it here
# would make a legacy deck disappear from the corpus without a word. It is scanned so
# that parse_file can reject it by name with a remedy the user can act on.
_SCANNED_SUFFIXES = {".pdf", ".pptx", ".ppt", ".docx"}


@dataclass
class TextBlock:
    block_type: Literal["text", "figure"]
    text: str
    # PDF page / PPTX slide, 0-indexed. DOCX: 1-based BLOCK ORDINAL, because a .docx
    # has no pages at all — see _parse_docx for why that is not papered over.
    page: int
    max_font_size: float            # max span font size; 0.0 for figure blocks
    span_font_sizes: list[float]    # all individual span sizes; empty for figures
    source_file: str                # filename only — not full path (S4, node_id stability)
    # "The format told us this is a heading" — PPTX placeholder type TITLE /
    # CENTER_TITLE, or a DOCX paragraph style named "Heading …" — as opposed to the
    # inferred relative-font-size path structure.py runs for PDF, where no such
    # declaration exists.
    is_explicit_heading: bool = False
    # Content flags. See the module docstring for detector-vs-heuristic status.
    has_table: bool = False
    has_code: bool = False
    has_equation: bool = False


@dataclass
class ParsedDocument:
    source_file: str                # filename only
    source_type: Literal["pdf", "pptx", "docx"]
    blocks: list[TextBlock]
    # PDF pages / PPTX slides. DOCX: the number of blocks, since a .docx has no pages.
    page_count: int
    warnings: list[str]             # pages with no text layer, size warnings, etc.


def parse_file(path: Path) -> ParsedDocument:
    """
    Parse one PDF, PPTX or DOCX file. Raises ValueError on any hard failure
    so the caller (parse_directory) can isolate it (R3).
    """
    _check_file_size(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".ppt":
        # Legacy PowerPoint is a binary OLE2 compound file, NOT an OOXML zip, so
        # python-pptx cannot read it under any circumstances. Handing it to
        # _parse_pptx produces an opaque zip error that reads like a corrupt file and
        # sends the user looking for damage that is not there. Name the format and
        # the remedy instead; parse_directory's per-file try/except (R3) carries this
        # message to the log and continues with the rest of the batch.
        raise ValueError(
            f"{path.name}: legacy .ppt is not supported (it is a binary format, "
            "not OOXML, and python-pptx cannot read it); convert it to .pptx and "
            "re-run."
        )
    raise ValueError(f"Unsupported file type: {suffix!r} ({path.name})")


def parse_directory(source_dir: Path) -> list[ParsedDocument]:
    """
    Parse all PDF/PPTX/DOCX files in source_dir.
    One bad file logs an error and is skipped; the batch always completes (R3).

    Each file is parsed under a config.PARSE_TIMEOUT_SECONDS watchdog (S3), so one
    pathological document cannot hang the whole ingest. A file that overruns is
    logged as failed and skipped, exactly like one that raises.

    LIMITATION — this is batch isolation, NOT preemption. `signal.alarm` is
    Unix-only, so the watchdog is a `ThreadPoolExecutor.result(timeout=...)`, and a
    Python thread cannot be forcibly killed. On timeout the batch moves on, but the
    abandoned parse keeps running in the background, still holding its CPU and memory;
    because executor workers are non-daemon threads it can also delay interpreter
    shutdown at the end of the run. What the requirement buys is that one bad file
    never blocks the others — not that its cost stops being paid. `multiprocessing`
    would give real preemption and is deliberately not used: it collides with
    Qdrant's exclusive file lock (S8) and adds complexity the project prohibits.
    """
    candidates = sorted(
        p for p in source_dir.iterdir()
        if p.is_file() and p.suffix.lower() in _SCANNED_SUFFIXES
    )
    if not candidates:
        logger.warning("No supported files found in %s", source_dir)

    docs: list[ParsedDocument] = []
    for path in candidates:
        try:
            doc = _parse_file_with_timeout(path)
            docs.append(doc)
            # page_count is not a page count for DOCX (there are no pages), so it is
            # reported by name rather than as "N pages".
            logger.info(
                "Parsed %s [%s]: %d blocks, page_count=%d",
                path.name, doc.source_type, len(doc.blocks), doc.page_count,
            )
        except Exception as exc:
            logger.error("Failed to parse %s — skipping: %s", path.name, exc)

    return docs


def _parse_file_with_timeout(path: Path) -> ParsedDocument:
    """
    Run `parse_file(path)` under the S3 per-file watchdog.

    Raises TimeoutError once the parse overruns config.PARSE_TIMEOUT_SECONDS. The
    worker thread is left running — see the limitation note in `parse_directory`.
    A fresh single-worker executor per file means the abandoned thread never starves
    the next file of a worker slot.
    """
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="parse")
    try:
        future = executor.submit(parse_file, path)
        try:
            return future.result(timeout=config.PARSE_TIMEOUT_SECONDS)
        except FuturesTimeoutError as exc:
            raise TimeoutError(
                f"{path.name}: parse exceeded PARSE_TIMEOUT_SECONDS="
                f"{config.PARSE_TIMEOUT_SECONDS}; abandoned (the worker thread cannot "
                "be killed and continues until process exit)"
            ) from exc
    finally:
        # wait=False: never block the batch on the thread we just abandoned.
        executor.shutdown(wait=False)


# ---------------------------------------------------------------------------
# PDF via PyMuPDF
# ---------------------------------------------------------------------------

def _parse_pdf(path: Path) -> ParsedDocument:
    import fitz  # PyMuPDF

    source_file = path.name
    blocks: list[TextBlock] = []
    warnings: list[str] = []

    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        raise ValueError(f"PyMuPDF could not open {path.name}: {exc}") from exc

    if doc.page_count > config.MAX_PAGES:
        doc.close()
        raise ValueError(
            f"{path.name}: {doc.page_count} pages exceeds MAX_PAGES={config.MAX_PAGES}"
        )

    for page_idx in range(doc.page_count):
        page = doc[page_idx]
        table_bboxes = _find_table_bboxes(page, page_idx, path.name, warnings)
        page_dict = page.get_text("dict")
        page_has_content = False

        for raw_block in page_dict.get("blocks", []):
            if raw_block["type"] == 1:                   # image block
                blocks.append(TextBlock(
                    block_type="figure",
                    text="",
                    page=page_idx,
                    max_font_size=0.0,
                    span_font_sizes=[],
                    source_file=source_file,
                ))
                page_has_content = True

            elif raw_block["type"] == 0:                 # text block
                span_texts: list[str] = []
                span_sizes: list[float] = []
                span_fonts: list[str] = []
                for line in raw_block.get("lines", []):
                    for span in line.get("spans", []):
                        t = span.get("text", "").strip()
                        if t:
                            span_texts.append(t)
                            size = span.get("size", 0.0)
                            if size > 0:
                                span_sizes.append(size)
                            font = span.get("font", "")
                            if font:
                                span_fonts.append(font)

                if not span_texts:
                    continue

                text = " ".join(span_texts)
                max_size = max(span_sizes) if span_sizes else 0.0
                blocks.append(TextBlock(
                    block_type="text",
                    text=text,
                    page=page_idx,
                    max_font_size=max_size,
                    span_font_sizes=span_sizes,
                    source_file=source_file,
                    has_table=_bbox_intersects_any(
                        raw_block.get("bbox"), table_bboxes
                    ),
                    has_code=is_code_font(span_fonts),
                    has_equation=is_equation_content(span_fonts, text),
                ))
                page_has_content = True

        if not page_has_content:
            # R10: log the cut, never silently drop content.
            msg = (
                f"Page {page_idx + 1} of {path.name} has no text or image layer. "
                "OCR is not available in MVP1 — page skipped."
            )
            logger.warning(msg)
            warnings.append(msg)

    page_count = doc.page_count
    doc.close()
    return ParsedDocument(
        source_file=source_file,
        source_type="pdf",
        blocks=blocks,
        page_count=page_count,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# PPTX via python-pptx
# ---------------------------------------------------------------------------

# Table cell/row joiners, shared by the PPTX and DOCX paths. Cell boundaries are
# preserved rather than flattened to spaces because the chunk text is what an item is
# later grounded in — "QuickSort | O(n log n)" keeps the row's association readable,
# "QuickSort O(n log n)" does not. Text, not markup: S5 requires model-facing content
# to stay inert.
_TABLE_CELL_SEP = " | "
_TABLE_ROW_SEP = "\n"


def _parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    _check_ooxml_zip_size(path)

    source_file = path.name
    blocks: list[TextBlock] = []
    warnings: list[str] = []

    _HEADING_TYPES = {
        PP_PLACEHOLDER.TITLE,         # 1
        PP_PLACEHOLDER.CENTER_TITLE,  # 3
    }

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ValueError(f"python-pptx could not open {path.name}: {exc}") from exc

    slide_count = len(prs.slides)
    if slide_count > config.MAX_PAGES:
        raise ValueError(
            f"{path.name}: {slide_count} slides exceeds MAX_PAGES={config.MAX_PAGES}"
        )

    for slide_idx, slide in enumerate(prs.slides):
        slide_has_content = False

        for shape in slide.shapes:
            if shape.has_table:
                # Exact, not a heuristic: python-pptx models a table shape directly.
                table = shape.table
                table_rows: list[str] = []
                cell_fonts: list[str] = []

                for r_idx in range(len(table.rows)):
                    cells: list[str] = []
                    for c_idx in range(len(table.columns)):
                        cell = table.cell(r_idx, c_idx)
                        # Skip positions covered by a merge. python-pptx puts the
                        # merged text on the origin cell and returns "" at every
                        # spanned position, so iterating rows/cells naively emits a
                        # stray blank cell for each one.
                        if cell.is_spanned:
                            continue
                        # Collapse intra-cell newlines so _TABLE_ROW_SEP stays
                        # meaningful as the row boundary.
                        cell_text = " ".join(cell.text.split())
                        if cell_text:
                            cells.append(cell_text)
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                font_name = run.font.name
                                if font_name:
                                    cell_fonts.append(font_name)
                    if cells:
                        table_rows.append(_TABLE_CELL_SEP.join(cells))

                table_text = _TABLE_ROW_SEP.join(table_rows)
                if not table_text:
                    warnings.append(
                        f"Slide {slide_idx + 1} of {path.name} has a table with no "
                        f"cell text — has_table is set but the node carries no table "
                        f"content to ground a question in."
                    )
                blocks.append(TextBlock(
                    block_type="text",
                    text=table_text,
                    page=slide_idx,
                    max_font_size=0.0,
                    span_font_sizes=[],
                    source_file=source_file,
                    has_table=True,
                    has_code=is_code_font(cell_fonts),
                    has_equation=is_equation_content(cell_fonts, table_text),
                ))
                slide_has_content = True
                continue

            if not shape.has_text_frame:
                if (
                    hasattr(shape, "shape_type")
                    and shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                ):
                    blocks.append(TextBlock(
                        block_type="figure",
                        text="",
                        page=slide_idx,
                        max_font_size=0.0,
                        span_font_sizes=[],
                        source_file=source_file,
                    ))
                    slide_has_content = True
                continue

            parts: list[str] = []
            run_fonts: list[str] = []
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    t = run.text.strip()
                    if t:
                        parts.append(t)
                        # None whenever the run inherits its face from the layout or
                        # theme, which is the common case — the font-name prong of the
                        # code/equation heuristics is simply blind on those runs.
                        font_name = run.font.name
                        if font_name:
                            run_fonts.append(font_name)

            if not parts:
                continue

            is_heading = (
                shape.is_placeholder
                and shape.placeholder_format is not None
                and shape.placeholder_format.type in _HEADING_TYPES
            )
            shape_text = " ".join(parts)
            blocks.append(TextBlock(
                block_type="text",
                text=shape_text,
                page=slide_idx,
                max_font_size=0.0,      # not used for PPTX heading detection
                span_font_sizes=[],
                source_file=source_file,
                is_explicit_heading=is_heading,
                has_code=is_code_font(run_fonts),
                has_equation=is_equation_content(run_fonts, shape_text),
            ))
            slide_has_content = True

        if not slide_has_content:
            warnings.append(f"Slide {slide_idx + 1} of {path.name} has no content.")

    return ParsedDocument(
        source_file=source_file,
        source_type="pptx",
        blocks=blocks,
        page_count=slide_count,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# DOCX via python-docx
# ---------------------------------------------------------------------------

def _parse_docx(path: Path) -> ParsedDocument:
    """
    Parse one .docx into TextBlocks.

    Headings are EXACT here, better than either other format: Word stores a real
    paragraph style, so a style named "Heading …" is the document itself declaring
    the paragraph a heading. `is_explicit_heading` is set from it and structure.py
    never runs the PDF font-size inference over a DOCX.

    ⚠ `page` IS A BLOCK ORDINAL, NOT A PAGE NUMBER. A .docx has no pages: pagination
    does not exist until Word lays the document out against a printer, the style
    definitions and the installed font metrics, and python-docx cannot compute it.
    Every block therefore carries its 1-based ordinal within the document. This is a
    disclosed limitation, not an approximation waiting to be improved: counting
    explicit page breaks would report "page 1" for the great majority of real
    documents, which contain none, and a plausible-looking wrong page number in a
    citation is worse than an honest ordinal. pipeline.md states the consequence for
    citations; todo.md carries the renderer-side fix (label a DOCX locator "¶12",
    not "p.12"), which needs the source type available at citation time.

    MAX_PAGES is deliberately NOT applied: there are no pages to count, and applying
    it to the block count would assert exactly the block==page equivalence this
    function exists to deny — a 2 000-block cap would also reject a legitimate
    hundred-page document. The S3 guards that do apply are MAX_FILE_SIZE_BYTES,
    MAX_DECOMPRESSED_SIZE_BYTES and the per-file PARSE_TIMEOUT_SECONDS watchdog.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_ooxml_zip_size(path)

    source_file = path.name
    blocks: list[TextBlock] = []
    warnings: list[str] = []

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ValueError(f"python-docx could not open {path.name}: {exc}") from exc

    p_tag, tbl_tag = qn("w:p"), qn("w:tbl")
    figures_by_paragraph = _docx_figures_by_paragraph(document, p_tag)
    figures_expected = len(document.inline_shapes)
    figures_placed = 0

    # Walk the body in document ORDER. `document.paragraphs` and `document.tables`
    # are two separate flat lists, so neither one alone can tell which paragraphs
    # follow which table — and section boundaries depend on that order.
    for child in document.element.body.iterchildren():
        if child.tag == tbl_tag:
            table_text, cell_fonts = _docx_table_text(Table(child, document))
            if not table_text:
                warnings.append(
                    f"Block {len(blocks) + 1} of {path.name} is a table with no cell "
                    f"text — has_table is set but the node carries no table content "
                    f"to ground a question in."
                )
            blocks.append(TextBlock(
                block_type="text",
                text=table_text,
                page=len(blocks) + 1,
                max_font_size=0.0,      # DOCX headings come from styles, not sizes
                span_font_sizes=[],
                source_file=source_file,
                has_table=True,         # exact: a w:tbl element is a table
                has_code=is_code_font(cell_fonts),
                has_equation=is_equation_content(cell_fonts, table_text),
            ))
            continue

        if child.tag != p_tag:
            continue                    # sectPr, bookmarks — no content of their own

        para = Paragraph(child, document)
        # Normalise whitespace: a paragraph may carry tabs and soft breaks, and
        # _TABLE_ROW_SEP only stays meaningful if a stray newline cannot appear in a
        # non-table block.
        text = " ".join(para.text.split())
        run_fonts = [r.font.name for r in para.runs if r.font.name]

        if text:
            blocks.append(TextBlock(
                block_type="text",
                text=text,
                page=len(blocks) + 1,
                max_font_size=0.0,
                span_font_sizes=[],
                source_file=source_file,
                is_explicit_heading=_is_docx_heading(para),
                has_code=is_code_font(run_fonts),
                has_equation=is_equation_content(run_fonts, text),
            ))

        for _ in range(figures_by_paragraph.get(child, 0)):
            blocks.append(TextBlock(
                block_type="figure",
                text="",
                page=len(blocks) + 1,
                max_font_size=0.0,
                span_font_sizes=[],
                source_file=source_file,
            ))
            figures_placed += 1

    if figures_placed != figures_expected:
        # R10: never silently drop content. An inline shape anchored somewhere the
        # body walk does not visit (a header, a footnote, a table cell) would
        # otherwise under-report has_figure with no trace.
        warnings.append(
            f"{path.name}: {figures_expected} inline shapes reported by python-docx "
            f"but {figures_placed} placed as figure blocks — has_figure may be "
            f"under-reported."
        )

    if not blocks:
        warnings.append(f"{path.name} has no paragraph, table or image content.")

    return ParsedDocument(
        source_file=source_file,
        source_type="docx",
        blocks=blocks,
        page_count=len(blocks),   # block count — a .docx has no pages
        warnings=warnings,
    )


def _is_docx_heading(paragraph: object) -> bool:
    """
    EXACT, not a heuristic: a paragraph style named with
    config.DOCX_HEADING_STYLE_PREFIX ("Heading 1" … "Heading 9") is the format
    declaring this paragraph a heading. No font-size inference is used for DOCX.
    """
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", None) or ""
    return name.startswith(config.DOCX_HEADING_STYLE_PREFIX)


def _docx_figures_by_paragraph(document: object, p_tag: str) -> dict[object, int]:
    """
    Map each of `document.inline_shapes` to the `w:p` element that contains it.

    `inline_shapes` is the authoritative list of embedded images but carries no
    position, and a figure block placed in the wrong section flags the wrong node —
    `requires_flags_any: ["has_figure"]` would then pick a section that has no
    figure. Walking up from the shape's `wp:inline` element restores the position.
    The caller cross-checks placed against expected and records a warning on any
    mismatch rather than quietly losing a shape.
    """
    counts: dict[object, int] = {}
    for shape in document.inline_shapes:
        node = shape._inline        # no public accessor for a shape's own element
        while node is not None and node.tag != p_tag:
            node = node.getparent()
        if node is not None:
            counts[node] = counts.get(node, 0) + 1
    return counts


def _docx_table_text(table: object) -> tuple[str, list[str]]:
    """
    Row-major text of a DOCX table plus the run font names inside it, joined with the
    same _TABLE_CELL_SEP / _TABLE_ROW_SEP the PPTX path uses so a row's associations
    survive into the grounding span.

    Merged cells behave the OPPOSITE way to python-pptx: python-docx resolves every
    spanned position to the SAME `w:tc` element and repeats its text there, where
    python-pptx returns "". A naive rows × columns walk therefore prints the merged
    text once per spanned position instead of emitting a blank. Cells are
    de-duplicated by `w:tc` identity across the whole table, which covers horizontal
    and vertical merges alike. The seen-list holds the element proxies, which keeps
    lxml's proxy identity stable for the `is` comparison.
    """
    rows_out: list[str] = []
    fonts: list[str] = []
    seen_tc: list[object] = []

    for r_idx in range(len(table.rows)):
        cells: list[str] = []
        for c_idx in range(len(table.columns)):
            cell = table.cell(r_idx, c_idx)
            tc = cell._tc          # no public accessor for a cell's own element
            if any(tc is seen for seen in seen_tc):
                continue
            seen_tc.append(tc)
            # Collapse intra-cell newlines so _TABLE_ROW_SEP stays meaningful as the
            # row boundary.
            cell_text = " ".join(cell.text.split())
            if cell_text:
                cells.append(cell_text)
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts.append(run.font.name)
        if cells:
            rows_out.append(_TABLE_CELL_SEP.join(cells))

    return _TABLE_ROW_SEP.join(rows_out), fonts


# ---------------------------------------------------------------------------
# Content flag heuristics
#
# Public because the tests exercise them directly: the font-name prong of
# has_equation cannot be driven end-to-end through a synthetic PDF (PyMuPDF's
# base-14 fonts substitute U+00B7 for every mathematical glyph and cannot embed a
# CMMI/Cambria-Math face without an external font file), so the unit is the seam
# that gets tested.
# ---------------------------------------------------------------------------

def is_code_font(font_names: list[str]) -> bool:
    """
    HEURISTIC: True if any font name looks monospace.

    Monospace is the whole signal — there is no lexical or syntactic analysis here
    and nothing in the pinned stack provides one. Prose set in a monospace face is a
    false positive; code set in a proportional face is a false negative.
    """
    return any(
        sub in name.lower()
        for name in font_names
        for sub in config.MONOSPACE_FONT_SUBSTRINGS
    )


def is_math_font(font_names: list[str]) -> bool:
    """HEURISTIC: True if any font name is a dedicated mathematics family."""
    return any(
        sub in name.lower()
        for name in font_names
        for sub in config.MATH_FONT_SUBSTRINGS
    )


def count_math_chars(text: str) -> int:
    """Number of characters in *text* drawn from config.MATH_UNICODE_RANGES."""
    total = 0
    for ch in text:
        cp = ord(ch)
        if any(lo <= cp <= hi for lo, hi in config.MATH_UNICODE_RANGES):
            total += 1
    return total


def is_equation_content(font_names: list[str], text: str) -> bool:
    """
    HEURISTIC: True if the block is set in a mathematics font, OR carries at least
    config.EQUATION_MIN_MATH_CHARS characters from config.MATH_UNICODE_RANGES.

    Neither prong sees mathematics that was typeset as an image, and MVP1 has no
    OCR — a scanned or picture-exported equation is a false negative.
    """
    if is_math_font(font_names):
        return True
    return count_math_chars(text) >= config.EQUATION_MIN_MATH_CHARS


def _find_table_bboxes(
    page: object,
    page_idx: int,
    file_name: str,
    warnings: list[str],
) -> list[tuple[float, float, float, float]]:
    """
    Bounding boxes of the tables PyMuPDF detects on *page*.

    Table detection is layout analysis over a third-party parse and can raise on
    unusual page content. That must not lose a whole 200-page document, so a failure
    degrades this one page to "no tables" and is recorded in ParsedDocument.warnings
    — visible in the ingest report, never swallowed.
    """
    try:
        finder = page.find_tables()
    except Exception as exc:
        msg = (
            f"Table detection failed on page {page_idx + 1} of {file_name}: {exc}. "
            "has_table may be under-reported for this page."
        )
        logger.warning(msg)
        warnings.append(msg)
        return []
    return [tuple(t.bbox) for t in finder.tables]


def _bbox_intersects_any(
    bbox: object,
    others: list[tuple[float, float, float, float]],
) -> bool:
    """True if *bbox* overlaps any rectangle in *others*. Both are (x0, y0, x1, y1)."""
    if not bbox or not others:
        return False
    ax0, ay0, ax1, ay1 = bbox
    for bx0, by0, bx1, by1 in others:
        if ax0 < bx1 and bx0 < ax1 and ay0 < by1 and by0 < ay1:
            return True
    return False


# ---------------------------------------------------------------------------
# Security checks (S3)
# ---------------------------------------------------------------------------

def _check_file_size(path: Path) -> None:
    size = path.stat().st_size
    if size > config.MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"{path.name}: {size} bytes exceeds MAX_FILE_SIZE_BYTES={config.MAX_FILE_SIZE_BYTES}"
        )


def _check_ooxml_zip_size(path: Path) -> None:
    """
    Zip-bomb guard for any OOXML container. PPTX and DOCX are both zips, and the
    decompression ratio risk is identical, so the check is shared rather than
    duplicated (it was `_check_pptx_zip_size` while PPTX was the only zip format).
    """
    try:
        with zipfile.ZipFile(path) as zf:
            total = sum(info.file_size for info in zf.infolist())
        if total > config.MAX_DECOMPRESSED_SIZE_BYTES:
            raise ValueError(
                f"{path.name}: decompressed {total} bytes exceeds "
                f"MAX_DECOMPRESSED_SIZE_BYTES={config.MAX_DECOMPRESSED_SIZE_BYTES}"
            )
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{path.name} is not a valid OOXML zip: {exc}") from exc
