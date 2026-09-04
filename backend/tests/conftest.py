"""Pytest configuration. The mocked LLM client is the default for all tests.
Live API tests require --live and are excluded from the default run (L3).
"""
from __future__ import annotations

from pathlib import Path

import pytest
from unittest.mock import MagicMock

from coursegen.llm.client import LLMClient, TokenBudget


def pytest_addoption(parser: pytest.Parser) -> None:
    parser.addoption(
        "--live",
        action="store_true",
        default=False,
        help="Run tests that make real LLM API calls (excluded by default).",
    )


def pytest_collection_modifyitems(
    config: pytest.Config, items: list[pytest.Item]
) -> None:
    if not config.getoption("--live"):
        skip_live = pytest.mark.skip(reason="Pass --live to run live API tests.")
        for item in items:
            if "live" in item.keywords:
                item.add_marker(skip_live)


@pytest.fixture
def mock_llm_client() -> MagicMock:
    """Default mocked LLM client. Replays a minimal valid response."""
    client = MagicMock(spec=LLMClient)
    client.budget = TokenBudget()
    client.call.return_value = {"mock": True, "estimated_tokens": 100}
    return client


@pytest.fixture
def token_budget() -> TokenBudget:
    return TokenBudget()


# ---------------------------------------------------------------------------
# P1 document fixtures — created once per test session via PyMuPDF / python-pptx.
# Using session scope so the heavy PDF creation is not repeated per test.
# ---------------------------------------------------------------------------

@pytest.fixture(scope="session")
def native_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    Native-text PDF: body at 11 pt, chapter headings at 24 pt.
    Font ratio is large enough that the relative threshold definitely fires.
    """
    import fitz  # PyMuPDF

    path = tmp_path_factory.mktemp("pdfs") / "native.pdf"
    doc = fitz.open()

    # Page 1 — Chapter 1
    p = doc.new_page()
    p.insert_text((72, 100), "Chapter 1: Introduction", fontsize=24)
    for y, line in enumerate([
        "This section introduces fundamental concepts.",
        "It covers the basic principles and key definitions.",
        "Understanding these concepts is essential for further study.",
        "We will explore various aspects in structured detail.",
        "Each concept builds carefully on the previous one.",
    ], start=1):
        p.insert_text((72, 100 + y * 60), line, fontsize=11)

    # Page 2 — Chapter 2
    p = doc.new_page()
    p.insert_text((72, 100), "Chapter 2: Data Structures", fontsize=24)
    for y, line in enumerate([
        "Data structures organise information efficiently in memory.",
        "Arrays provide fast indexed access to homogeneous elements.",
        "Linked lists allow dynamic allocation without contiguous memory.",
        "Trees represent hierarchical parent-child relationships.",
    ], start=1):
        p.insert_text((72, 100 + y * 60), line, fontsize=11)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def slide_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    Slide-exported PDF: body at 24 pt, slide titles at 36 pt.
    Regression fixture: an absolute font-size threshold would classify
    every 24 pt body block as a heading. The relative threshold must not.
    Gate: heading_fraction < HEADING_MAX_BLOCK_PCT (0.30).
    """
    import fitz

    path = tmp_path_factory.mktemp("pdfs") / "slides.pdf"
    doc = fitz.open()

    for i in range(5):
        p = doc.new_page()
        p.insert_text((72, 80), f"Lecture {i + 1}: Module Overview", fontsize=36)
        p.insert_text((72, 220), "First key point with supporting detail.", fontsize=24)
        p.insert_text((72, 320), "Second key point elaborated with an example.", fontsize=24)
        p.insert_text((72, 420), "Third key point summarising the concept.", fontsize=24)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def figure_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """PDF that contains at least one image block so has_figure can be verified."""
    import fitz

    path = tmp_path_factory.mktemp("pdfs") / "figure.pdf"
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Section with Diagram", fontsize=20)
    p.insert_text((72, 200), "Figure description text follows below.", fontsize=11)

    # Insert a 4×4 white rectangle as an image using PyMuPDF's own Pixmap.
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4))
    pix.clear_with(255)  # white
    img_rect = fitz.Rect(72, 250, 200, 350)
    p.insert_image(img_rect, pixmap=pix)

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def malformed_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Garbage bytes with a .pdf extension — must fail gracefully."""
    path = tmp_path_factory.mktemp("pdfs") / "malformed.pdf"
    path.write_bytes(b"This is not a valid PDF file, just garbage bytes.")
    return path


@pytest.fixture(scope="session")
def legacy_ppt(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    A file carrying the legacy `.ppt` extension.

    Its contents are deliberately irrelevant: `.ppt` is a binary OLE2 compound file,
    not an OOXML zip, so python-pptx can never read one. The parser must reject it on
    the suffix with an actionable message rather than letting python-pptx fail
    opaquely — which reads like a corrupt file and sends the user hunting for damage
    that is not there. The bytes below are the real OLE2 magic number so the fixture
    is not merely garbage.
    """
    path = tmp_path_factory.mktemp("legacy") / "old_deck.ppt"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    return path


@pytest.fixture(scope="session")
def minimal_pptx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Two-slide PPTX with title placeholders."""
    from pptx import Presentation

    path = tmp_path_factory.mktemp("pptx") / "test.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Introduction to Algorithms"
    slide.placeholders[1].text = "This slide covers fundamental algorithm concepts."

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Sorting Algorithms"
    slide.placeholders[1].text = "We discuss merge sort and quicksort in detail."

    prs.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Content-flag fixtures — one document per detectable flag, plus a corpus that
# combines them. These exist because P1 shipped with has_table / has_equation /
# has_code hardcoded False and no fixture could have noticed.
# ---------------------------------------------------------------------------

def _tiny_png_bytes() -> bytes:
    """A 4x4 white PNG, produced by PyMuPDF so no image library is needed."""
    import fitz

    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4))
    pix.clear_with(255)
    return pix.tobytes("png")


def _draw_ruled_table(page, x0: float, y0: float, rows: int, cols: int) -> None:
    """Draw a ruled grid with text in every cell — what page.find_tables() looks for."""
    import fitz

    cell_w, cell_h = 110.0, 28.0
    for r in range(rows + 1):
        page.draw_line(
            fitz.Point(x0, y0 + r * cell_h),
            fitz.Point(x0 + cols * cell_w, y0 + r * cell_h),
        )
    for c in range(cols + 1):
        page.draw_line(
            fitz.Point(x0 + c * cell_w, y0),
            fitz.Point(x0 + c * cell_w, y0 + rows * cell_h),
        )
    for r in range(rows):
        for c in range(cols):
            page.insert_text(
                (x0 + c * cell_w + 4, y0 + r * cell_h + 19),
                f"cell {r}{c}",
                fontsize=10,
            )


@pytest.fixture(scope="session")
def table_code_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    PDF carrying a ruled table (page 1) and a Courier code listing (page 2).

    Table: detected by page.find_tables(). Code: detected by the monospace-font-name
    heuristic — PyMuPDF's base-14 "cour" reports font name "Courier".
    """
    import fitz

    path = tmp_path_factory.mktemp("flagpdf") / "table_code.pdf"
    doc = fitz.open()

    p = doc.new_page()
    p.insert_text((72, 60), "Chapter 1: Comparison Table", fontsize=24)
    _draw_ruled_table(p, 72, 100, rows=3, cols=3)

    p = doc.new_page()
    p.insert_text((72, 60), "Chapter 2: Reference Implementation", fontsize=24)
    for i, line in enumerate([
        "def quicksort(items):",
        "    if len(items) <= 1:",
        "        return items",
        "    pivot = items[0]",
    ], start=1):
        p.insert_text((72, 100 + i * 24), line, fontsize=11, fontname="cour")

    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture(scope="session")
def math_pptx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    PPTX carrying mathematics as Unicode text, a monospace run, a real table shape
    and a picture shape.

    The equation document is a PPTX and not a PDF for a concrete reason: PyMuPDF's
    base-14 fonts substitute U+00B7 for every mathematical glyph, so a synthetic PDF
    physically cannot carry U+2200-range characters in its text layer, and no math
    font file may be added (no new dependencies). PPTX stores run text as XML and
    round-trips them exactly. The PDF side of the same heuristic is covered by the
    unit tests on `is_equation_content`.
    """
    from pptx import Presentation
    from pptx.util import Inches

    tmp_dir = tmp_path_factory.mktemp("flagpptx")
    path = tmp_dir / "math.pptx"
    png_path = tmp_dir / "dot.png"
    png_path.write_bytes(_tiny_png_bytes())

    prs = Presentation()
    content_layout = prs.slide_layouts[1]   # Title and Content
    title_only_layout = prs.slide_layouts[5]

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Gradient Descent Update"
    slide.placeholders[1].text_frame.text = (
        "For all ∀x ∈ S the update ∂L/∂w ≠ 0 converges "
        "when ∑ error ≤ ∞."
    )

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Training Loop"
    body = slide.placeholders[1].text_frame
    body.text = "for epoch in range(n): step(model, batch)"
    body.paragraphs[0].runs[0].font.name = "Consolas"

    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "Benchmark Results"
    bench = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    ).table
    bench.cell(0, 0).text = "Model"
    bench.cell(0, 1).text = "Accuracy"
    bench.cell(1, 0).text = "Baseline"
    bench.cell(1, 1).text = "0.71"

    slide = prs.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "Architecture Diagram"
    slide.shapes.add_picture(str(png_path), Inches(1), Inches(2), Inches(1), Inches(1))

    prs.save(str(path))
    return path


@pytest.fixture(scope="session")
def structured_docx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    DOCX carrying every signal `_parse_docx` reads: real "Heading 1" / "Heading 2"
    paragraph styles, body prose, a populated table, a Consolas run, Unicode
    mathematics and an inline image.

    Heading styles are the point of the fixture. DOCX is the only one of the three
    formats where a heading is a FACT rather than an inference — `paragraph.style.name`
    is set by Word — so the test asserts exact heading text, not a fraction below a
    threshold as the PDF path must.
    """
    from docx import Document
    from docx.shared import Inches

    tmp_dir = tmp_path_factory.mktemp("docx")
    path = tmp_dir / "structured.docx"
    png_path = tmp_dir / "dot.png"
    png_path.write_bytes(_tiny_png_bytes())

    document = Document()
    document.add_heading("Chapter 1: Sorting Algorithms", level=1)
    document.add_paragraph(
        "Sorting arranges elements into a defined order and underpins searching, "
        "deduplication and a great many database operations."
    )

    document.add_heading("1.1 QuickSort", level=2)
    code_run = document.add_paragraph().add_run(
        "def quicksort(items): return items"
    )
    code_run.font.name = "Consolas"
    document.add_paragraph(
        "For all ∀x ∈ S the recurrence ∑ T(n) ≤ ∞ bounds "
        "the expected work."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Algorithm"
    table.cell(0, 1).text = "Complexity"
    table.cell(1, 0).text = "QuickSort"
    table.cell(1, 1).text = "O(n log n)"

    document.add_heading("1.2 Architecture Diagram", level=2)
    document.add_paragraph().add_run().add_picture(str(png_path), width=Inches(1))

    document.save(str(path))
    png_path.unlink()   # keep only the parseable document in the fixture directory
    return path


@pytest.fixture(scope="session")
def rich_source_dir(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """
    A synthetic course corpus exercising every flag the pinned stack can detect.

    Deliberately larger than the other fixtures: it is the input to the ingest →
    solver integration test, which needs enough leaf sections to give
    `final_default.json` (32 slots) a realistic span budget, rather than starving on
    corpus size and hiding what the flag distribution is really doing.

    Contents: a 6-page PDF of 18 chapters (2 with embedded figures, 1 with a ruled
    table, 1 with a Courier listing), plus a 14-slide PPTX (3 with Unicode
    mathematics, 1 table shape, 1 monospace run, 1 picture).
    """
    import fitz
    from pptx import Presentation
    from pptx.util import Inches

    source_dir = tmp_path_factory.mktemp("rich_corpus")

    # --- PDF: 6 pages x 3 chapters ---
    doc = fitz.open()
    chapter = 0
    for _page_idx in range(6):
        page = doc.new_page()
        for slot in range(3):
            chapter += 1
            top = 60 + slot * 230
            page.insert_text((72, top), f"Chapter {chapter}: Topic {chapter}", fontsize=24)
            for line_idx, line in enumerate([
                f"Topic {chapter} introduces its core definitions and vocabulary.",
                f"The material for topic {chapter} builds on the preceding chapter.",
                f"Worked reasoning for topic {chapter} is developed step by step.",
            ], start=1):
                page.insert_text((72, top + line_idx * 26), line, fontsize=11)

            if chapter == 4:                      # ruled table
                _draw_ruled_table(page, 72, top + 100, rows=3, cols=3)
            elif chapter == 8:                    # monospace listing
                page.insert_text(
                    (72, top + 110), "while queue: node = queue.pop(0)",
                    fontsize=11, fontname="cour",
                )
            elif chapter in (12, 16):             # embedded figures
                pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4))
                pix.clear_with(255)
                page.insert_image(
                    fitz.Rect(72, top + 105, 172, top + 165), pixmap=pix
                )

    doc.save(str(source_dir / "lecture_notes.pdf"))
    doc.close()

    # --- PPTX: 14 slides ---
    png_path = source_dir / "_dot.png"
    png_path.write_bytes(_tiny_png_bytes())

    prs = Presentation()
    content_layout = prs.slide_layouts[1]
    title_only_layout = prs.slide_layouts[5]

    math_bodies = [
        "Because ∀x ∈ D we have ∑ residual ≤ ∞ at convergence.",
        "The bound ∂J/∂θ ≠ 0 holds whenever ∏ terms ≥ 1.",
        "Note ∫ density ≡ 1 and ∑ weights ≤ ∞ for every layer.",
    ]
    for idx in range(14):
        if idx < 3:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = f"Derivation {idx + 1}"
            slide.placeholders[1].text_frame.text = math_bodies[idx]
        elif idx == 3:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = "Reference Loop"
            body = slide.placeholders[1].text_frame
            body.text = "for row in table: emit(row)"
            body.paragraphs[0].runs[0].font.name = "Consolas"
        elif idx == 4:
            slide = prs.slides.add_slide(title_only_layout)
            slide.shapes.title.text = "Measured Results"
            results = slide.shapes.add_table(
                2, 3, Inches(1), Inches(2), Inches(5), Inches(1)
            ).table
            for col, header in enumerate(("Method", "Latency", "Throughput")):
                results.cell(0, col).text = header
            for col, value in enumerate(("Batched", "42 ms", "980 rps")):
                results.cell(1, col).text = value
        elif idx == 5:
            slide = prs.slides.add_slide(title_only_layout)
            slide.shapes.title.text = "System Diagram"
            slide.shapes.add_picture(
                str(png_path), Inches(1), Inches(2), Inches(1), Inches(1)
            )
        else:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = f"Lecture Topic {idx + 1}"
            slide.placeholders[1].text_frame.text = (
                f"Slide {idx + 1} explains its topic with supporting detail and "
                "a worked example for the reader to follow."
            )

    prs.save(str(source_dir / "slides.pptx"))
    png_path.unlink()   # keep only the two parseable documents in the corpus
    return source_dir


@pytest.fixture(autouse=True, scope="session")
def _enable_test_seams():
    """Turn on the `/api/internal/*` routes for the suite.

    They are 404 by default so a running server does not expose an
    unauthenticated way to clear the consent gate (`API-CONTRACT.md` has always
    said "Not for the UI"; nothing enforced it). The tests that reset disclosure
    between cases need them, and this is the only place that should switch them
    on.
    """
    import os

    previous = os.environ.get("PREPIFY_TEST_SEAMS")
    os.environ["PREPIFY_TEST_SEAMS"] = "1"
    yield
    if previous is None:
        os.environ.pop("PREPIFY_TEST_SEAMS", None)
    else:
        os.environ["PREPIFY_TEST_SEAMS"] = previous
